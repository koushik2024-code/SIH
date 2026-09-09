import io
import json
import socket
from pathlib import Path
from urllib.parse import urlsplit, parse_qs
import pytest
from app.config import settings, ASSETS
from app.db import db
from app import llm, vector_store
from app.security import make_citation_token, make_token
from app.repository import source_evidence
from app.schemas import OUTPUT_MODELS
from app.verify import verify
from app.ingest import parse_file, canonical_text
from app.chunking import chunk_sections
from app.router import choose_route


def valid_output(kind, eid):
    cite = [eid]
    if kind == 'linkedin':
        return {'hook': 'A new approach to translation', 'body': ['The research uses attention to process sequences.'], 'takeaway': 'Review the reported findings.', 'hashtags': ['#Research'], 'citations': cite}
    if kind == 'executive_summary':
        return {'title': 'Research brief', 'summary': 'The team reports a Transformer architecture.', 'key_points': ['Attention processes sequences.'], 'citations': cite}
    if kind == 'x_post':
        return {'posts': ['The team reports a Transformer architecture using attention.'], 'citations': cite}
    if kind == 'email':
        return {'subject': 'Research update', 'greeting': 'Hello,', 'body': ['The team reports a Transformer architecture.'], 'call_to_action': '', 'sign_off': 'Regards', 'citations': cite}
    if kind == 'advisory':
        return {'title': 'Research information', 'summary': 'The team reports a Transformer architecture.', 'affected': [], 'sections': [], 'recommendations': [], 'citations': cite}
    if kind == 'infographic':
        return {'title': 'Transformer research', 'headline': 'Attention for translation', 'key_points': ['Attention processes sequences.'], 'visual_elements': ['Attention diagram'], 'layout': 'Vertical factual panels', 'citations': cite}
    if kind == 'presentation':
        return {'slides': [{'title': 'Research', 'bullets': ['Attention processes sequences.'], 'speaker_notes': 'Review the research.', 'citations': cite}]}
    return {'title': 'Research briefing', 'scenes': [{'scene': i, 'narration': 'The team reports a Transformer architecture.', 'caption': 'Attention and translation', 'visual_keywords': ['research'], 'citations': cite} for i in range(1, 5)]}


def fake_model(monkeypatch, kind, eid):
    monkeypatch.setattr(llm, 'ollama_json', lambda messages, schema: json.dumps(valid_output(kind, eid)))


def test_login_and_token_type_separation(client, source):
    good = client.post('/api/auth/login', json={'username': 'alice', 'password': 'test-pass'})
    assert good.status_code == 200
    assert client.post('/api/auth/login', json={'username': 'alice', 'password': 'wrong'}).status_code == 401
    citation = make_citation_token(client.users['alice'], f'EVID-{source}-0000')
    assert client.get('/api/sources', headers={'Authorization': 'Bearer ' + citation}).status_code == 401
    assert client.get('/api/assets').status_code == 401


def test_direct_and_immutable_history(client, source, monkeypatch):
    eid = f'EVID-{source}-0000'
    fake_model(monkeypatch, 'linkedin', eid)
    body = {'source_ids': [source], 'output_types': ['linkedin'], 'retrieval_query': 'must be ignored'}
    first = client.post('/api/transform', headers=client.headers_for('alice'), json=body)
    assert first.status_code == 200, first.text
    a = first.json()['outputs'][0]
    assert a['route'] == 'DIRECT' and a['retrieval_query_used'] is None
    assert a['verification']['factual_accuracy_verified'] is False
    second = client.post('/api/transform', headers=client.headers_for('alice'), json=body).json()['outputs'][0]
    assert a['asset_id'] != second['asset_id']
    with db() as con:
        paths = [r['path'] for r in con.execute('SELECT path FROM assets').fetchall()]
    assert len(set(paths)) == 2 and all(Path(p).is_file() for p in paths)
    assert len(client.get('/api/assets', headers=client.headers_for('alice')).json()) == 2
    export = client.get(a['download_url'], headers=client.headers_for('alice'))
    assert export.status_code == 200 and 'A new approach' in export.text and '/api/citations/' in export.text


def test_cross_user_source_asset_and_route(client, source, monkeypatch):
    fake_model(monkeypatch, 'linkedin', f'EVID-{source}-0000')
    result = client.post('/api/transform', headers=client.headers_for('alice'), json={'source_ids':[source], 'output_types':['linkedin']}).json()
    aid = result['outputs'][0]['asset_id']
    h = client.headers_for('bob')
    for path in [f'/api/sources/{source}', f'/api/assets/{aid}', f'/api/assets/{aid}/download']:
        assert client.get(path, headers=h).status_code == 404
    for path, body in [('/api/route', {'source_ids':[source]}), ('/api/transform', {'source_ids':[source], 'output_types':['linkedin']})]:
        assert client.post(path, headers=h, json=body).status_code == 404
    assert client.get('/api/assets', headers=h).json() == []


def test_signed_citations_expiry_mismatch_and_ownership(client, source):
    eid = f'EVID-{source}-0000'
    uid = client.users['alice']
    valid = make_citation_token(uid, eid)
    assert client.get(f'/api/citations/{eid}', params={'token': valid}).status_code == 200
    assert client.get('/api/citations/EVID-999-0000', params={'token': valid}).status_code == 403
    assert client.get(f'/api/citations/{eid}', params={'token': make_citation_token(uid, eid, -1)}).status_code == 401
    assert client.get(f'/api/citations/{eid}', params={'token': make_citation_token(client.users['bob'], eid)}).status_code == 404
    assert client.get(f'/api/citations/{eid}', params={'token': make_token(uid, 'alice')}).status_code == 401
    page = client.get(f'/api/citations/{eid}', params={'token': valid})
    assert page.headers['referrer-policy'] == 'no-referrer' and page.headers['cache-control'] == 'no-store'


def test_repair_once_and_invalid_citations_never_saved(client, source, monkeypatch):
    eid = f'EVID-{source}-0000'
    replies = iter(['{}', json.dumps(valid_output('linkedin', eid))])
    calls = []
    def model(messages, schema):
        calls.append(messages.copy())
        return next(replies)
    monkeypatch.setattr(llm, 'ollama_json', model)
    result = client.post('/api/transform', headers=client.headers_for('alice'), json={'source_ids':[source], 'output_types':['linkedin']})
    assert result.status_code == 200 and len(calls) == 2
    assert 'failed validation' in calls[1][-1]['content']
    bad = valid_output('linkedin', eid); bad['citations'] = ['Vaswani et al. 2017']
    monkeypatch.setattr(llm, 'ollama_json', lambda *args: json.dumps(bad))
    result = client.post('/api/transform', headers=client.headers_for('alice'), json={'source_ids':[source], 'output_types':['linkedin']})
    assert result.status_code == 422
    assert len(client.get('/api/assets', headers=client.headers_for('alice')).json()) == 1


@pytest.mark.parametrize('kind', list(OUTPUT_MODELS))
def test_output_contracts(kind):
    eid = 'EVID-1-0000'
    value = valid_output(kind, eid)
    evidence = [{'evidence_id':eid,'content':'The research reports attention.'}]
    assert verify(kind, value, evidence)['passed']
    value['unexpected_format_field'] = []
    assert not verify(kind, value, evidence)['passed']


def test_empty_length_and_severity():
    evidence = [{'evidence_id':'EVID-1-0000','content':'The system uses attention.'}]
    assert not verify('linkedin', {}, evidence)['passed']
    value = valid_output('linkedin','EVID-1-0000'); value['body'] = ['  ']
    assert not verify('linkedin', value, evidence)['passed']
    value = valid_output('x_post','EVID-1-0000'); value['posts'] = ['x' * 281]
    assert not verify('x_post', value, evidence)['passed']
    value = valid_output('advisory','EVID-1-0000'); value['severity'] = 'Critical'
    assert not verify('advisory', value, evidence)['passed']


def test_route_preview_large_and_explicit(client, source):
    h = client.headers_for('alice')
    assert client.post('/api/route', headers=h, json={'source_ids':[source]}).json()['route'] == 'DIRECT'
    big = client.post('/api/sources/text', headers=h, json={'text':'Satellite policy. Financial risks. Cybersecurity recommendations.\n\n' * 700}).json()['source_id']
    decision = client.post('/api/route', headers=h, json={'source_ids':[big]}).json()
    assert decision['route'] == 'RAG' and decision['query_recommended']
    assert client.post('/api/route', headers=h, json={'source_ids':[source],'persistent_knowledge':True}).json()['route'] == 'RAG'
    assert choose_route('tiny', 6).route == 'RAG'


def test_real_qdrant_tenant_filters_and_query(client, source, monkeypatch):
    # Actual embedded Qdrant; synthetic vectors remove model/network dependence.
    bob = client.post('/api/sources/text', headers=client.headers_for('bob'), json={'text':'Bob private translation material'}).json()['source_id']
    monkeypatch.setattr(vector_store, 'embed_documents', lambda texts: [[1.0, 0.0, 0.0] for _ in texts])
    queries = []
    def embed(query):
        queries.append(query)
        return [1.0, 0.0, 0.0]
    monkeypatch.setattr(vector_store, 'embed_query', embed)
    vector_store.index_source(client.users['bob'], bob)
    vector_store.index_source(client.users['alice'], source)
    hits = vector_store.search_evidence(client.users['alice'], 'translation findings', [source], 10)
    assert hits and all(e['user_id'] == client.users['alice'] for e in hits)
    fake_model(monkeypatch, 'linkedin', f'EVID-{source}-0000')
    result = client.post('/api/transform', headers=client.headers_for('alice'), json={'source_ids':[source], 'output_types':['linkedin'], 'retrieval_requested':True, 'retrieval_query':'attention methodology only'})
    assert result.status_code == 200, result.text
    assert queries[-1] == 'attention methodology only'
    assert result.json()['outputs'][0]['retrieval_query_used'] == queries[-1]
    vector_store.delete_source_vectors(client.users['alice'], source)
    assert vector_store.search_evidence(client.users['alice'], 'x', [source], 10) == []
    assert vector_store.search_evidence(client.users['bob'], 'x', [bob], 10)


def test_multi_source_rag_and_partial_success(client, source, monkeypatch):
    other = client.post('/api/sources/text', headers=client.headers_for('alice'), json={'text':'Additional research findings on attention.'}).json()['source_id']
    monkeypatch.setattr(vector_store, 'embed_documents', lambda texts: [[1.0, 0.0, 0.0] for _ in texts])
    monkeypatch.setattr(vector_store, 'embed_query', lambda query: [1.0, 0.0, 0.0])
    fake_model(monkeypatch, 'linkedin', f'EVID-{source}-0000')
    result = client.post('/api/transform', headers=client.headers_for('alice'), json={'source_ids':[source,other], 'output_types':['linkedin','email'], 'retrieval_requested':True})
    assert result.status_code == 200
    assert len(result.json()['outputs']) == 1 and result.json()['errors'][0]['output_type'] == 'email'


@pytest.mark.parametrize('url', ['http://127.0.0.1', 'http://[::1]', 'http://169.254.169.254', 'http://10.0.0.1', 'http://user:pass@example.com', 'file:///etc/passwd', 'http://localhost:8000'])
def test_ssrf_rejections(url):
    from app.utils import validate_public_url
    with pytest.raises(ValueError):
        validate_public_url(url)


def test_office_and_structured_chunks(tmp_path):
    from docx import Document
    from openpyxl import Workbook
    doc = Document(); doc.add_heading('Research',1); doc.add_paragraph('The research reports an attention architecture.'); table = doc.add_table(rows=1,cols=2); table.cell(0,0).text='Metric';table.cell(0,1).text='Value'
    path = tmp_path/'research.docx'; doc.save(path)
    text,meta,sections = parse_file(path)
    assert meta['tables'] == 1 and sections[1]['heading'] == 'Research'
    workbook = Workbook(); workbook.active.title='Revenue';workbook.active.append(['Year','Amount']);workbook.active.append([2026,100]);workbook.create_sheet('Risks').append(['Risk','Impact'])
    path = tmp_path/'data.xlsx';workbook.save(path)
    text,meta,sections = parse_file(path)
    chunks = chunk_sections(sections, 1, 1, path.name)
    assert meta['sheets'] == ['Revenue','Risks'] and chunks[0]['sheet'] == 'Revenue'
    assert chunks[0]['row_start'] == 2 and 'Year: 2026' in chunks[0]['content']


def test_pdf_direct_summary(client, monkeypatch):
    from reportlab.pdfgen import canvas
    data = io.BytesIO(); page = canvas.Canvas(data); page.drawString(60,700,'The team reports a Transformer architecture using attention for translation.');page.save()
    response = client.post('/api/sources/upload',headers=client.headers_for('alice'),files={'file':('small.pdf',data.getvalue(),'application/pdf')})
    assert response.status_code == 200, response.text
    sid = response.json()['source_id']; fake_model(monkeypatch,'executive_summary',f'EVID-{sid}-0000')
    output = client.post('/api/transform',headers=client.headers_for('alice'),json={'source_ids':[sid],'output_types':['executive_summary']})
    assert output.status_code == 200 and output.json()['route'] == 'DIRECT'
    assert source_evidence(client.users['alice'],[sid])[0]['page'] == 1


def test_frontend_and_input_errors(client):
    assert client.get('/').status_code == 200
    assert client.get('/app.js').status_code == 200
    assert client.get('/style.css').status_code == 200
    h = client.headers_for('alice')
    assert client.post('/api/route', headers=h, json={'source_ids':[]}).status_code == 422
    assert client.post('/api/sources/upload', headers=h, files={'file':('bad.exe',b'payload')}).status_code == 400
    assert client.post('/api/sources/text', headers=h, json={'text':'   '}).status_code == 422


def test_ollama_failure_502(client, source, monkeypatch):
    def fail(*args):
        raise llm.ModelUnavailable('Local Ollama unavailable')
    monkeypatch.setattr(llm,'ollama_json',fail)
    response = client.post('/api/transform',headers=client.headers_for('alice'),json={'source_ids':[source],'output_types':['linkedin']})
    assert response.status_code == 502


def test_real_image_and_scanned_pdf_ocr(client, tmp_path, monkeypatch):
    import shutil
    from PIL import Image, ImageDraw, ImageFont
    if not shutil.which('tesseract'):
        pytest.skip('Tesseract not installed')
    image = Image.new('RGB',(1500,400),'white')
    font = ImageFont.truetype('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf',36) if Path('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf').exists() else ImageFont.truetype('arial.ttf',36)
    ImageDraw.Draw(image).text((50,100),'Research uses attention for translation.',font=font,fill='black')
    for suffix, kind in [('.png','email'),('.pdf','executive_summary')]:
        path = tmp_path / ('scan' + suffix); image.save(path)
        response = client.post('/api/sources/upload',headers=client.headers_for('alice'),files={'file':(path.name,path.read_bytes())})
        assert response.status_code == 200, response.text
        sid = response.json()['source_id']
        ev = source_evidence(client.users['alice'],[sid])
        assert ev[0]['ocr'] and 'attention' in ev[0]['content'].lower()
        if suffix == '.pdf':
            assert response.json()['metadata']['ocr_pages'] == [1]
        fake_model(monkeypatch, kind, ev[0]['evidence_id'])
        result = client.post('/api/transform',headers=client.headers_for('alice'),json={'source_ids':[sid],'output_types':[kind]})
        assert result.status_code == 200


def test_audio_video_ingest_adapter(client, tmp_path, monkeypatch):
    import wave, subprocess, shutil
    from types import SimpleNamespace
    from app import ingest
    if not shutil.which('ffmpeg'):
        pytest.skip('FFmpeg not installed')
    wav = tmp_path/'sample.wav'
    with wave.open(str(wav),'wb') as output:
        output.setnchannels(1);output.setsampwidth(2);output.setframerate(16000);output.writeframes(b'\x00\x00'*16000)
    class SpeechFixture:
        def transcribe(self,path,**kwargs):
            return iter([SimpleNamespace(text='The research uses attention.',start=0.0,end=1.0)]), SimpleNamespace(duration=1.0,language='en')
    monkeypatch.setattr(ingest,'whisper_model',lambda:SpeechFixture())
    video = tmp_path/'input.mp4'
    subprocess.run(['ffmpeg','-v','error','-y','-f','lavfi','-i','color=c=blue:s=320x180:d=1','-i',str(wav),'-c:v','libx264','-c:a','aac','-shortest',str(video)],check=True,capture_output=True)
    for path in [wav,video]:
        result = client.post('/api/sources/upload',headers=client.headers_for('alice'),files={'file':(path.name,path.read_bytes())})
        assert result.status_code == 200,result.text
        evidence = source_evidence(client.users['alice'],[result.json()['source_id']])
        assert evidence[0]['timestamp_start'] == 0 and evidence[0]['timestamp_end'] == 1
        fake_model(monkeypatch,'executive_summary',evidence[0]['evidence_id'])
        generated = client.post('/api/transform',headers=client.headers_for('alice'),json={'source_ids':[result.json()['source_id']],'output_types':['executive_summary']})
        assert generated.status_code == 200


def test_deterministic_video_real_ffmpeg(client, source, monkeypatch, tmp_path):
    import wave, shutil, subprocess
    from app import render
    if not shutil.which('ffmpeg'):
        pytest.skip('FFmpeg not installed')
    voice = tmp_path / 'fixture.onnx'
    voice.write_bytes(b'fixture-not-a-real-model')
    voice.with_suffix('.onnx.json').write_text('{"language":{"code":"en_US"}}')
    monkeypatch.setattr(settings, 'piper_model', str(voice))
    monkeypatch.setattr(settings, 'enable_web_images', False)
    durations = iter([0.72,0.93,1.12,0.81])
    def narration_fixture(text, path):
        with wave.open(str(path),'wb') as audio:
            audio.setnchannels(1);audio.setsampwidth(2);audio.setframerate(24000)
            audio.writeframes(b'\x00\x00'*int(next(durations)*24000))
    monkeypatch.setattr(render,'_tts',narration_fixture)
    fake_model(monkeypatch,'video',f'EVID-{source}-0000')
    response = client.post('/api/transform',headers=client.headers_for('alice'),json={'source_ids':[source],'output_types':['video']})
    assert response.status_code == 200,response.text
    output = response.json()['outputs'][0]
    assert [round(s['duration_seconds'],2) for s in output['media']['scenes']] == [0.72,0.93,1.12,0.81]
    with db() as con:
        path = con.execute('SELECT path FROM assets WHERE id=?',(output['asset_id'],)).fetchone()['path']
    probe = subprocess.run(['ffprobe','-v','error','-show_entries','format=duration','-of','json',path],check=True,capture_output=True,text=True)
    duration = float(json.loads(probe.stdout)['format']['duration'])
    assert abs(duration - sum([.72,.93,1.12,.81])) < .25
    downloaded = client.get(output['download_url'],headers=client.headers_for('alice'))
    assert downloaded.status_code == 200 and downloaded.headers['content-type'] == 'video/mp4'
    import io, zipfile
    from PIL import Image
    frames = client.get(f"/api/assets/{output['asset_id']}/frames", headers=client.headers_for('alice'))
    assert frames.status_code == 200
    with zipfile.ZipFile(io.BytesIO(frames.content)) as archive:
        assert len([n for n in archive.namelist() if n.endswith('.png')]) == 4
        assert 'timeline.json' in archive.namelist()
    assert client.get(f"/api/assets/{output['asset_id']}/frames", headers=client.headers_for('bob')).status_code == 404
    extracted = tmp_path / 'decoded.png'
    subprocess.run(['ffmpeg','-v','error','-y','-ss','0.3','-i',path,'-frames:v','1','-update','1',str(extracted)],check=True)
    with Image.open(extracted) as image:
        assert image.size == (1280,720)
        assert image.getpixel((980,190)) != image.getpixel((0,0))



def test_url_ip_pinning_redirect_and_size(monkeypatch):
    from app import utils
    import urllib3
    monkeypatch.setattr(socket,'getaddrinfo',lambda *args,**kwargs:[(socket.AF_INET,socket.SOCK_STREAM,6,'',('93.184.216.34',443))])
    captured = {}
    class Reply:
        status=200
        headers={'Content-Type':'text/html'}
        def stream(self,*args,**kwargs):
            yield b'<html>Research</html>'
        def close(self):
            pass
    class Pool:
        def __init__(self,host,**kwargs):
            captured.update(host=host,**kwargs)
        def urlopen(self,method,path,**kwargs):
            captured.update(method=method,path=path,**kwargs)
            return Reply()
        def close(self):
            pass
    monkeypatch.setattr(urllib3,'HTTPSConnectionPool',Pool)
    data,_ = utils.fetch_public_url('https://example.com/article')
    assert b'Research' in data and captured['host'] == '93.184.216.34' and captured['server_hostname'] == 'example.com'
    assert captured['redirect'] is False and captured['headers']['Host'] == 'example.com'
    Reply.status = 302
    with pytest.raises(ValueError,match='Redirects'):
        utils.fetch_public_url('https://example.com/article')
    Reply.status = 200
    monkeypatch.setattr(settings,'max_url_mb',0)
    with pytest.raises(OverflowError):
        utils.fetch_public_url('https://example.com/article')


def test_upload_stream_limit(client, monkeypatch):
    monkeypatch.setattr(settings,'max_upload_mb',1)
    response = client.post('/api/sources/upload',headers=client.headers_for('alice'),files={'file':('oversized.txt',b'x'*(3*1024*1024))})
    assert response.status_code == 413,response.text


@pytest.mark.parametrize('kind', [k for k in OUTPUT_MODELS if k != 'video'])
def test_every_text_output_renders_and_downloads(client, source, monkeypatch, kind):
    fake_model(monkeypatch,kind,f'EVID-{source}-0000')
    response=client.post('/api/transform',headers=client.headers_for('alice'),json={'source_ids':[source],'output_types':[kind]})
    assert response.status_code == 200,response.text
    output=response.json()['outputs'][0]
    assert output['rendered_text'].strip()
    assert client.get(output['download_url'],headers=client.headers_for('alice')).status_code == 200


def test_large_pdf_routes_through_real_qdrant(client, monkeypatch):
    from reportlab.pdfgen import canvas
    data=io.BytesIO();pdf=canvas.Canvas(data)
    for page in range(35):
        for row in range(20):
            pdf.drawString(40,750-row*30,'Research uses attention for translation. Satellite policy and financial risks are additional topics.')
        pdf.showPage()
    pdf.save()
    response=client.post('/api/sources/upload',headers=client.headers_for('alice'),files={'file':('large_topics.pdf',data.getvalue(),'application/pdf')})
    assert response.status_code == 200,response.text
    sid=response.json()['source_id']
    route=client.post('/api/route',headers=client.headers_for('alice'),json={'source_ids':[sid]}).json()
    assert route['route'] == 'RAG' and route['query_recommended']
    monkeypatch.setattr(vector_store,'embed_documents',lambda texts:[[1.,0.,0.] for _ in texts])
    monkeypatch.setattr(vector_store,'embed_query',lambda query:[1.,0.,0.])
    def reply(messages,schema):
        import re
        # Cite an ID actually supplied in retrieved context, whatever tie order Qdrant uses.
        eid=re.search(r'EVID-\d+-\d+',messages[1]['content']).group()
        return json.dumps(valid_output('linkedin',eid))
    monkeypatch.setattr(llm,'ollama_json',reply)
    generated=client.post('/api/transform',headers=client.headers_for('alice'),json={'source_ids':[sid],'output_types':['linkedin'],'retrieval_query':'attention and translation'})
    assert generated.status_code == 200,generated.text
    assert generated.json()['outputs'][0]['retrieval_query_used'] == 'attention and translation'


def test_create_user_module_command(tmp_path):
    import os, subprocess, sys, sqlite3
    env=dict(os.environ,DATA_DIR=str(tmp_path/'module-data'))
    result=subprocess.run([sys.executable,'-m','scripts.create_user','demo','xyz'],env=env,capture_output=True,text=True)
    assert result.returncode == 0,result.stderr
    with sqlite3.connect(tmp_path/'module-data'/'app.db') as con:
        row=con.execute('SELECT username,password_hash FROM users').fetchone()
    assert row[0] == 'demo' and row[1].startswith('$argon2')
